"""Excel-Export: Kostenschaetzung und Auffaelligkeiten.

Blaetter:
- Zusammenfassung:     Total exkl. MwSt., Quelldateien, Datum, Kennzahlen
- Kosten je Position:  eine Zeile pro Pos-ID (mit eBKP-T-Code) und Totalzeile
- Bauteilliste:        eine Zeile pro Bauteil (GUID, Fachmodell, Pos-ID, Menge, Mengenquelle, Kosten)
- Auffaelligkeiten:    alle Pruefmeldungen aus allen Schritten, Fehler zuerst

Die Datei wird entweder gespeichert (ziel = Pfad) oder als Bytes zurueckgegeben
(ziel = None), z.B. fuer den Download-Knopf in Streamlit.
"""

from datetime import datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from validation.model_checks import FEHLER

# Spaltenname im Programm -> Spaltenkopf im Excel; nur diese Spalten werden exportiert
SPALTEN_KOSTEN = {
    "pos_id": "Pos-ID",
    "ebkp_code": "eBKP-T Code",
    "gruppe": "Gruppe",
    "bezeichnung": "Bezeichnung",
    "mengenart": "Mengenart",
    "anzahl_bauteile": "Anzahl Bauteile",
    "menge": "Menge",
    "einheit": "Einheit",
    "einheitspreis": "Einheitspreis (CHF)",
    "kosten": "Kosten (CHF)",
    "anteil": "Anteil",
    "hinweis": "Hinweis",
}

SPALTEN_BAUTEILE = {
    "guid": "GUID",
    "fachmodell": "Fachmodell",
    "datei": "Datei",
    "klasse": "IFC-Klasse",
    "name": "Name",
    "bereich": "Bereich",
    "bauteilgruppe": "Bauteilgruppe",
    "typ_schicht": "Typ / Schicht",
    "element": "Element",
    "pos_id": "Pos-ID",
    "treffer_regeln": "Regel (Zeile Zuordnung)",
    "menge": "Menge",
    "einheit": "Einheit",
    "mengenquelle": "Mengenquelle",
    "ersatzfelder": "Aus Ersatzquelle",
    "einheitspreis": "Einheitspreis (CHF)",
    "kosten": "Kosten (CHF)",
}

SPALTEN_MELDUNGEN = {
    "schwere": "Schwere",
    "schritt": "Schritt",
    "pruefung": "Prüfung",
    "guid": "GUID",
    "fachmodell": "Fachmodell",
    "blatt": "Blatt",
    "excel_zeile": "Zeile",
    "pos_id": "Pos-ID",
    "meldung": "Meldung",
}

# Zahlenformate je Excel-Spaltenkopf
FORMATE = {
    "Menge": "#,##0.000",
    "Einheitspreis (CHF)": "#,##0.00",
    "Kosten (CHF)": "#,##0.00",
    "Anteil": "0.0%",
}

KOPF_SCHRIFT = Font(bold=True, color="FFFFFF")
KOPF_FARBE = PatternFill("solid", fgColor="1F4E78")
FEHLER_FARBE = PatternFill("solid", fgColor="F8CBAD")
HINWEIS_FARBE = PatternFill("solid", fgColor="FFF2CC")
MAX_SPALTENBREITE = 60


def _waehle_spalten(tabelle, spalten):
    """Nimmt die gewuenschten Spalten (falls vorhanden) und gibt ihnen die Excel-Namen."""
    vorhanden = [s for s in spalten if s in tabelle.columns]
    return tabelle[vorhanden].rename(columns=spalten)


def _gestalte_blatt(blatt, kopfzeile=1):
    """Kopfzeile farbig, Zahlenformate, Spaltenbreiten, Filter und fixierte Kopfzeile."""
    for zelle in blatt[kopfzeile]:
        zelle.font = KOPF_SCHRIFT
        zelle.fill = KOPF_FARBE

    for spalte in blatt.iter_cols(min_row=kopfzeile):
        kopf = spalte[0].value
        format_ = FORMATE.get(kopf)
        if format_:
            for zelle in spalte[1:]:
                zelle.number_format = format_
        # Breite nach dem laengsten Text der Spalte
        laengste = max(len(str(zelle.value)) for zelle in spalte if zelle.value is not None)
        blatt.column_dimensions[get_column_letter(spalte[0].column)].width = min(laengste + 2, MAX_SPALTENBREITE)

    blatt.freeze_panes = blatt.cell(row=kopfzeile + 1, column=1)
    blatt.auto_filter.ref = f"A{kopfzeile}:{get_column_letter(blatt.max_column)}{blatt.max_row}"


def _faerbe_meldungen(blatt):
    """Faerbt Fehler rot und Hinweise gelb (Spalte A = Schwere)."""
    for zeile in blatt.iter_rows(min_row=2):
        farbe = FEHLER_FARBE if zeile[0].value == FEHLER else HINWEIS_FARBE
        for zelle in zeile:
            zelle.fill = farbe


def sammle_meldungen(meldungen_je_schritt):
    """Haengt die Meldungstabellen aller Schritte untereinander.

    meldungen_je_schritt: {"Modellprüfung": tabelle, "Listenprüfung": tabelle, ...}
    Die Tabellen haben unterschiedliche Spalten (GUID oder Excel-Zeile), fehlende bleiben leer.
    Fehler kommen zuerst.
    """
    tabellen = [t.assign(schritt=schritt) for schritt, t in meldungen_je_schritt.items() if len(t)]
    if not tabellen:
        return pd.DataFrame(columns=list(SPALTEN_MELDUNGEN))
    alle = pd.concat(tabellen, ignore_index=True)
    alle["ist_fehler"] = alle["schwere"] == FEHLER
    alle = alle.sort_values("ist_fehler", ascending=False, kind="stable").drop(columns="ist_fehler")
    return alle.reset_index(drop=True)


def _zusammenfassung(kosten, bauteile, meldungen, quelldateien, liste):
    """Kennzahlen als Tabelle mit zwei Spalten (Angabe, Wert)."""
    zugeordnet = bauteile["pos_id"].notna().sum()
    konflikte = bauteile["treffer_pos_ids"].str.contains(",", na=False).sum()
    zeilen = [
        ("Kostenschätzung KOSTRA", ""),
        ("Erstellt am", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Total exkl. MwSt. (CHF)", round(kosten["kosten"].sum(), 2)),
        ("", ""),
        ("Bauteile im Modell", len(bauteile)),
        ("davon einer Position zugeordnet", zugeordnet),
        ("davon ohne Position", len(bauteile) - zugeordnet - konflikte),
        ("davon mit Konflikt (mehrere Positionen)", konflikte),
        ("Positionen in der Liste", len(kosten)),
        ("davon ohne Bauteil im Modell", (kosten["anzahl_bauteile"] == 0).sum()),
        ("davon ohne Preis", kosten["einheitspreis"].isna().sum()),
        ("Fehler", (meldungen["schwere"] == FEHLER).sum()),
        ("Hinweise", (meldungen["schwere"] != FEHLER).sum()),
        ("", ""),
        ("eBKP-T-Liste", liste or ""),
    ]
    zeilen += [("Quelldatei" if i == 0 else "", Path(d).name) for i, d in enumerate(quelldateien)]
    return pd.DataFrame(zeilen, columns=["Angabe", "Wert"])


def schreibe_kostenschaetzung(kosten, bauteile, meldungen_je_schritt, quelldateien, liste=None, ziel=None):
    """Schreibt die Kostenschaetzung als Excel.

    kosten, bauteile:      Ergebnis von calculation.costs.berechne_kosten
    meldungen_je_schritt:  {"Modellprüfung": ..., "Listenprüfung": ..., "Zuordnung": ..., "Berechnung": ...}
    quelldateien:          Namen oder Pfade der IFC-Dateien
    liste:                 Name der eBKP-T-Liste
    ziel:                  Pfad zum Speichern; None -> die Datei wird als Bytes zurueckgegeben
    """
    meldungen = sammle_meldungen(meldungen_je_schritt)

    kosten_excel = _waehle_spalten(kosten, SPALTEN_KOSTEN)
    total = {"Bezeichnung": "Total exkl. MwSt.", "Anzahl Bauteile": kosten["anzahl_bauteile"].sum(),
             "Kosten (CHF)": kosten["kosten"].sum(), "Anteil": 1.0}
    kosten_excel = pd.concat([kosten_excel, pd.DataFrame([total])], ignore_index=True)

    blaetter = {
        "Zusammenfassung": _zusammenfassung(kosten, bauteile, meldungen, quelldateien, liste),
        "Kosten je Position": kosten_excel,
        "Bauteilliste": _waehle_spalten(bauteile, SPALTEN_BAUTEILE),
        "Auffälligkeiten": _waehle_spalten(meldungen, SPALTEN_MELDUNGEN),
    }

    datei = BytesIO() if ziel is None else ziel
    with pd.ExcelWriter(datei, engine="openpyxl") as writer:
        for name, tabelle in blaetter.items():
            tabelle.to_excel(writer, sheet_name=name, index=False)
            _gestalte_blatt(writer.sheets[name])

        # Einzelheiten: Totalzeilen fett, Total-Betrag formatiert, Meldungen eingefaerbt
        for zelle in writer.sheets["Kosten je Position"][writer.sheets["Kosten je Position"].max_row]:
            zelle.font = Font(bold=True)
        zusammenfassung = writer.sheets["Zusammenfassung"]
        zusammenfassung.auto_filter.ref = None
        zusammenfassung["A2"].font = Font(bold=True, size=14)
        zusammenfassung["B4"].number_format = "#,##0.00"
        zusammenfassung["A4"].font = zusammenfassung["B4"].font = Font(bold=True)
        _faerbe_meldungen(writer.sheets["Auffälligkeiten"])

    if ziel is None:
        return datei.getvalue()
    return None


def schreibe_auffaelligkeiten(meldungen_je_schritt, bauteile=None, ziel=None):
    """Nur die Meldungen als Excel, z.B. fuer den Download nach der Modellpruefung (Schritt 3).

    Mit bauteile werden zu jeder Meldung die Angaben des betroffenen Bauteils (per GUID) ergaenzt,
    damit man es im Modell findet und korrigieren kann.
    """
    meldungen = sammle_meldungen(meldungen_je_schritt)
    tabelle = _waehle_spalten(meldungen, SPALTEN_MELDUNGEN)
    # Nur Meldungen mit GUID betreffen ein Bauteil (Listenpruefung hat keine GUID-Spalte)
    if bauteile is not None and "guid" in meldungen.columns and meldungen["guid"].notna().any():
        angaben = ["guid", "klasse", "name", "bereich", "bauteilgruppe", "typ_schicht", "element", "datei"]
        details = meldungen[["guid"]].merge(bauteile[angaben].drop_duplicates("guid"), on="guid", how="left")
        details = _waehle_spalten(details.drop(columns="guid"), SPALTEN_BAUTEILE)
        tabelle = pd.concat([tabelle, details], axis=1)

    datei = BytesIO() if ziel is None else ziel
    with pd.ExcelWriter(datei, engine="openpyxl") as writer:
        tabelle.to_excel(writer, sheet_name="Auffälligkeiten", index=False)
        _gestalte_blatt(writer.sheets["Auffälligkeiten"])
        _faerbe_meldungen(writer.sheets["Auffälligkeiten"])

    if ziel is None:
        return datei.getvalue()
    return None


if __name__ == "__main__":
    # Test im Terminal: python -m export.excel_writer
    from calculation.costs import berechne_kosten
    from calculation.mapping import ordne_zu
    from importers.ebkp_reader import lese_ebkp_liste
    from importers.ifc_reader import lese_alle_ifc
    from validation.list_checks import pruefe_liste
    from validation.model_checks import pruefe_modell

    liste = Path("daten/eBKP-T_Preisliste_Zuordnung.xlsx")
    positionen, zuordnung, einstellungen = lese_ebkp_liste(liste)
    ifc_dateien = sorted(Path("daten").glob("*.ifc"))
    bauteile = lese_alle_ifc("daten", einstellungen)
    zugeordnet, zuordnungsmeldungen = ordne_zu(bauteile, zuordnung, positionen)
    kosten, bauteilkosten, mengenmeldungen = berechne_kosten(zugeordnet, positionen)

    meldungen = {
        "Modellprüfung": pruefe_modell(bauteile),
        "Listenprüfung": pruefe_liste(positionen, zuordnung, einstellungen),
        "Zuordnung": zuordnungsmeldungen,
        "Berechnung": mengenmeldungen,
    }
    ziel = Path("daten/Kostenschaetzung_Test.xlsx")   # in daten/, damit es nicht ins Git kommt
    schreibe_kostenschaetzung(kosten, bauteilkosten, meldungen, ifc_dateien, liste.name, ziel)
    print("Gespeichert:", ziel)

    # Kontrolle: Datei wieder einlesen
    for blatt, tabelle in pd.read_excel(ziel, sheet_name=None).items():
        print(f"  {blatt:20} {len(tabelle):4} Zeilen")
    print("Bytes-Variante fuer Streamlit:", len(schreibe_kostenschaetzung(kosten, bauteilkosten, meldungen, ifc_dateien)), "Bytes")
